from fastapi import FastAPI, UploadFile, File
from fastapi.responses import HTMLResponse, JSONResponse
from openpyxl import load_workbook
from io import BytesIO

app = FastAPI(title='Jalkheri Report Automation')

HTML = '''<!doctype html><html><head><meta name='viewport' content='width=device-width,initial-scale=1'><title>Jalkheri DCS Dashboard</title><style>body{font-family:Arial;margin:0;background:#f4f6f8;color:#17202a}.wrap{max-width:1200px;margin:auto;padding:28px}.card,.kpi{background:#fff;border-radius:14px;padding:20px;margin:18px 0;box-shadow:0 2px 10px #0001}.grid{display:grid;grid-template-columns:repeat(auto-fit,minmax(190px,1fr));gap:14px}button{background:#1769aa;color:#fff;border:0;padding:11px 18px;border-radius:8px;cursor:pointer}table{width:100%;border-collapse:collapse}td,th{padding:9px;border-bottom:1px solid #eee;text-align:left}input{padding:10px;border:1px solid #ccd3da;border-radius:8px}.value{font-size:25px;font-weight:700;margin-top:7px}.label{font-size:13px;color:#667085}</style></head><body><div class='wrap'><h1>Jalkheri DCS Report Automation</h1><p>Upload the latest Jalkheri Excel file.</p><div class='card'><input id='file' type='file' accept='.xlsx,.xlsm'> <button onclick='upload()'>Process Jalkheri</button><p id='status'></p></div><div id='dashboard' style='display:none'><div class='grid' id='cards'></div><div class='card'><h2>Jalkheri Tag Data</h2><input id='search' placeholder='Search tag or description' oninput='renderTable()'><div id='table'></div></div></div></div><script>let rows=[];async function upload(){let f=document.getElementById('file').files[0];if(!f)return alert('Select an Excel file first');let fd=new FormData();fd.append('file',f);document.getElementById('status').textContent='Processing...';let r=await fetch('/api/upload',{method:'POST',body:fd});let d=await r.json();if(!r.ok){document.getElementById('status').textContent=d.detail;return}rows=d.rows;document.getElementById('status').textContent='Processed '+d.count+' tags from '+d.sheet+'.';document.getElementById('dashboard').style.display='block';document.getElementById('cards').innerHTML=Object.entries(d.summary).map(x=>'<div class="kpi"><div class="label">'+x[0]+'</div><div class="value">'+x[1]+'</div></div>').join('');renderTable()}function renderTable(){let q=document.getElementById('search').value.toLowerCase();let x=rows.filter(r=>(r.tag+' '+r.description).toLowerCase().includes(q)).slice(0,300);document.getElementById('table').innerHTML='<table><tr><th>Tag</th><th>Description</th><th>Unit</th><th>Latest Value</th></tr>'+x.map(r=>'<tr><td>'+r.tag+'</td><td>'+r.description+'</td><td>'+r.unit+'</td><td>'+r.value+'</td></tr>').join('')+'</table>'}</script></body></html>'''

@app.get('/', response_class=HTMLResponse)
def home(): return HTML

def read_jalkheri(data):
    wb=load_workbook(BytesIO(data), data_only=True, read_only=True)
    if 'Jalkheri' not in wb.sheetnames: raise ValueError('Jalkheri sheet not found')
    ws=wb['Jalkheri']; last_col=0
    for c in range(7, ws.max_column+1):
        if any(ws.cell(r,c).value is not None for r in range(8,ws.max_row+1)): last_col=c
    rows=[]
    for r in range(8,ws.max_row+1):
        tag=ws.cell(r,3).value; desc=ws.cell(r,5).value; unit=ws.cell(r,6).value
        val=ws.cell(r,last_col).value if last_col else ws.cell(r,4).value
        if tag: rows.append({'tag':str(tag),'description':str(desc or ''),'unit':str(unit or ''),'value':val})
    return rows,ws.cell(5,5).value

@app.post('/api/upload')
async def upload(file: UploadFile=File(...)):
    try:
        rows,date=read_jalkheri(await file.read())
        def find(term):
            for r in rows:
                if term.lower() in (r['tag']+' '+r['description']).lower() and isinstance(r['value'],(int,float)): return r['value']
            return '—'
        return {'sheet':'Jalkheri','count':len(rows),'date':str(date) if date else '—','summary':{'TG Load (MW)':find('MW001'),'Main Steam Flow (TPH)':find('MAIN STM FLOW 1'),'Live Steam Pressure':find('LIVE STM PR'),'Main Steam Temperature':find('LIVE STM TEMP'),'Furnace Draft':find('FURNACE DRAFT')},'rows':rows}
    except Exception as e: return JSONResponse({'detail':str(e)},status_code=400)
